"""
eMAG 爬虫 - 保存模块

输出格式: CSV / Excel / JSON + 图片下载
"""

import csv
import json
import os
from pathlib import Path

import config
from logger import logger


# ============================================================
# 字段扁平化
# ============================================================
def _flatten_dict(d: dict, prefix: str = "") -> dict:
    """将嵌套字典扁平化为一层"""
    flat = {}
    for k, v in d.items():
        key = f"{prefix}{k}" if prefix else k
        if isinstance(v, dict):
            flat.update(_flatten_dict(v, f"{key}_"))
        elif isinstance(v, list):
            flat[key] = "|".join(str(x) for x in v)
        else:
            flat[key] = v
    return flat


def _normalize_value(v) -> str:
    """将任意值转为 CSV 安全字符串"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    # 移除换行符
    return str(v).replace("\n", " ").replace("\r", " ")


# ============================================================
# 收集所有字段名（从已规范化的产品中动态发现）
# ============================================================
def _collect_all_fields(products: list[dict]) -> list[str]:
    """收集所有产品中出现过的全部字段，保持首个产品的顺序"""
    seen = set()
    fields = []
    for p in products:
        for k in p:
            if k not in seen:
                seen.add(k)
                fields.append(k)
    return fields


# ============================================================
# CSV 写入
# ============================================================
def save_csv(products: list[dict], filepath: str = None) -> str:
    """保存为 CSV，字段自动展开"""
    from field_mapper import normalize_product

    filepath = filepath or config.CSV_FILE
    os.makedirs(os.path.dirname(filepath), exist_ok=True)

    if not products:
        logger.warning("没有产品数据可保存到 CSV")
        return filepath

    products = [normalize_product(p) for p in products]
    fields = _collect_all_fields(products)
    logger.info(f"CSV 字段总数: {len(fields)}")

    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()

        for p in products:
            row = {k: _normalize_value(p.get(k)) for k in fields}
            writer.writerow(row)

    logger.info(f"CSV 已保存: {filepath} ({len(products)} 行)")
    return filepath


# ============================================================
# Excel 写入
# ============================================================
def save_excel(products: list[dict], filepath: str = None) -> str:
    """保存为 Excel（需要 openpyxl）"""
    from field_mapper import normalize_product

    filepath = filepath or config.EXCEL_FILE
    os.makedirs(os.path.dirname(filepath), exist_ok=True)

    if not products:
        logger.warning("没有产品数据可保存到 Excel")
        return filepath

    products = [normalize_product(p) for p in products]

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("缺少 openpyxl，请执行: pip install openpyxl")
        return ""

    fields = _collect_all_fields(products)
    logger.info(f"Excel 字段总数: {len(fields)}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"

    # 标题行
    for col_idx, field in enumerate(fields, 1):
        ws.cell(row=1, column=col_idx, value=field)

    # 数据行
    for row_idx, p in enumerate(products, 2):
        for col_idx, field in enumerate(fields, 1):
            ws.cell(row=row_idx, column=col_idx, value=_normalize_value(p.get(field)))

    # 冻结首行
    ws.freeze_panes = "A2"

    wb.save(filepath)
    logger.info(f"Excel 已保存: {filepath} ({len(products)} 行)")
    return filepath


# ============================================================
# JSON 写入
# ============================================================
def save_json(products: list[dict], filepath: str = None) -> str:
    """保存为 JSON"""
    from field_mapper import normalize_product

    filepath = filepath or config.JSON_FILE
    os.makedirs(os.path.dirname(filepath), exist_ok=True)

    if not products:
        logger.warning("没有产品数据可保存到 JSON")
        return filepath

    products = [normalize_product(p) for p in products]

    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(products, f, ensure_ascii=False, indent=2, default=str)

    logger.info(f"JSON 已保存: {filepath} ({len(products)} 条)")
    return filepath


# ============================================================
# 图片下载（委托给 ImageDownloader）
# ============================================================
def download_all_images(products: list[dict], fetcher_cls) -> None:
    """
    下载所有商品的全部图库图片

    委托 image_handler.ImageDownloader 处理。
    """
    from image_handler import ImageDownloader

    if not config.DOWNLOAD_IMAGES:
        logger.info("图片下载已禁用")
        return

    # 对于没有 _gallery_urls 的商品（仅列表页数据），使用 image_url 作为兜底
    for p in products:
        if not p.get("_gallery_urls") and p.get("image_url"):
            from image_handler import get_high_res_url
            p["_gallery_urls"] = [get_high_res_url(p.get("image_url", ""))]

    downloader = ImageDownloader(fetcher_cls)
    downloader.download_all_products(products)
